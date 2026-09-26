"""Generate professional Aegis Guard pitch deck (PPTX) matching product brand."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

OUT = Path(__file__).resolve().parent.parent / "submission" / "Aegis_Guard_ET_AI_Hackathon_Problem2.pptx"

# Brand (from public/index.html)
INK = RGBColor(0x0B, 0x1F, 0x1A)
ACCENT = RGBColor(0x0F, 0x76, 0x6E)
ACCENT2 = RGBColor(0x11, 0x5E, 0x59)
MUTED = RGBColor(0x5B, 0x6F, 0x68)
BG = RGBColor(0xF4, 0xF7, 0xF6)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LINE = RGBColor(0xD9, 0xE2, 0xDE)
ALLOW = RGBColor(0x04, 0x78, 0x57)
BLOCK = RGBColor(0xB9, 0x1C, 0x1C)


def _set_run(run, text, size=18, bold=False, color=INK, font="Arial"):
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font


def _add_bg(slide, color=BG):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    # send to back
    spTree = slide.shapes._spTree
    sp = shape._element
    spTree.remove(sp)
    spTree.insert(2, sp)


def _bar(slide, y=0):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, y, Inches(13.333), Inches(0.08))
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT
    bar.line.fill.background()


def _footer(slide, page: str):
    box = slide.shapes.add_textbox(Inches(0.6), Inches(7.05), Inches(10), Inches(0.3))
    p = box.text_frame.paragraphs[0]
    run = p.add_run()
    _set_run(run, f"Aegis Guard  ·  ET AI Hackathon Problem 2  ·  Claim F3 / D2  ·  {page}", 10, False, MUTED)
    box2 = slide.shapes.add_textbox(Inches(11.2), Inches(7.05), Inches(1.5), Inches(0.3))
    p2 = box2.text_frame.paragraphs[0]
    p2.alignment = PP_ALIGN.RIGHT
    r2 = p2.add_run()
    _set_run(r2, page, 10, False, MUTED)


def _title(slide, text, top=0.45, size=32):
    box = slide.shapes.add_textbox(Inches(0.7), Inches(top), Inches(12), Inches(0.7))
    p = box.text_frame.paragraphs[0]
    run = p.add_run()
    _set_run(run, text, size, True, INK)


def _body(slide, lines, top=1.4, size=16, width=12):
    box = slide.shapes.add_textbox(Inches(0.7), Inches(top), Inches(width), Inches(5.2))
    tf = box.text_frame
    tf.word_wrap = True
    first = True
    for line in lines:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_after = Pt(8)
        run = p.add_run()
        _set_run(run, line, size, False, INK if not line.startswith("  ") else MUTED)


def _card(slide, left, top, w, h, title, body, accent=ACCENT):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = WHITE
    shape.line.color.rgb = LINE
    tip = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(0.08), Inches(h))
    tip.fill.solid()
    tip.fill.fore_color.rgb = accent
    tip.line.fill.background()
    t = slide.shapes.add_textbox(Inches(left + 0.25), Inches(top + 0.2), Inches(w - 0.4), Inches(0.4))
    r = t.text_frame.paragraphs[0].add_run()
    _set_run(r, title, 14, True, ACCENT2)
    b = slide.shapes.add_textbox(Inches(left + 0.25), Inches(top + 0.65), Inches(w - 0.4), Inches(h - 0.85))
    b.text_frame.word_wrap = True
    rb = b.text_frame.paragraphs[0].add_run()
    _set_run(rb, body, 12, False, MUTED)


def build() -> Path:
    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # 1 Title
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _add_bg(s, INK)
    accent_band = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(5.9), Inches(13.333), Inches(1.6))
    accent_band.fill.solid()
    accent_band.fill.fore_color.rgb = ACCENT2
    accent_band.line.fill.background()
    box = s.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(11.5), Inches(1))
    r = box.text_frame.paragraphs[0].add_run()
    _set_run(r, "AEGIS GUARD", 48, True, WHITE)
    box2 = s.shapes.add_textbox(Inches(0.8), Inches(2.5), Inches(11.5), Inches(1.2))
    tf = box2.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    _set_run(r, "Behavioral Twin Prompt Injection Firewall", 24, False, RGBColor(0x9A, 0xE6, 0xC5))
    p2 = tf.add_paragraph()
    r2 = p2.add_run()
    _set_run(r2, "ET AI Hackathon: Agentic Edition  ·  Problem 2  ·  Solo", 16, False, RGBColor(0xC5, 0xD2, 0xCC))
    box3 = s.shapes.add_textbox(Inches(0.8), Inches(6.2), Inches(11.5), Inches(0.9))
    tf3 = box3.text_frame
    r = tf3.paragraphs[0].add_run()
    _set_run(r, "Self-estimated 9-blocker:  F3  /  D2", 20, True, WHITE)
    p = tf3.add_paragraph()
    r = p.add_run()
    _set_run(r, "Live demo  ·  aegis-prompt-injection-firewall.vercel.app", 14, False, RGBColor(0xE6, 0xF5, 0xF2))

    # 2 Problem
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _add_bg(s)
    _bar(s)
    _title(s, "The enterprise risk")
    _card(s, 0.7, 1.4, 3.9, 2.4, "Agents touch untrusted content", "PDF, email, HTML, OCR, and APIs now feed tool-using agents. Attackers hide instructions inside those documents.")
    _card(s, 4.8, 1.4, 3.9, 2.4, "Indirect prompt injection", "A single poisoned attachment can trigger get_secret, send_email, or shell tools — without the user noticing.")
    _card(s, 8.9, 1.4, 3.7, 2.4, "Classifiers are not enough", "Keyword / LLM jailbreak detectors ask: “Does this look bad?” Attackers evade that every day.", BLOCK)
    _body(s, [
        "Accenture clients are shipping copilots and agents into production workflows.",
        "Without a runtime control plane, every document becomes an attack surface.",
    ], top=4.2, size=15)
    _footer(s, "02")

    # 3 Insight
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _add_bg(s)
    _bar(s)
    _title(s, "The insight: runtime, not classification")
    _body(s, [
        "Prompt injection is not a text-classification problem.",
        "",
        "It is a runtime security problem:",
        "    Does untrusted content change what tools the agent would call?",
        "",
        "Aegis Guard runs a behavioral twin:",
        "    • Trusted plan  — instructions with untrusted content stripped",
        "    • Raw plan      — instructions as attackers hope the agent will see them",
        "    • If plans diverge on high-risk tools → BLOCK or QUARANTINE",
        "",
        "Everyone else ships a classifier. We ship a twin.",
    ], top=1.35, size=17)
    _footer(s, "03")

    # 4 Solution
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _add_bg(s)
    _bar(s)
    _title(s, "Aegis Guard control plane")
    steps = [
        ("01 Ingest", "Normalize PDF, HTML, email, OCR, API"),
        ("02 Decode", "Unwrap base64 / hex / rot13 / homoglyphs"),
        ("03 Provenance", "Tag trusted vs untrusted spans"),
        ("04 Static", "Detect 9 attack families"),
        ("05 Twin", "Dual tool-plan comparison"),
        ("06 Risk + Gate", "ALLOW · QUARANTINE · BLOCK"),
    ]
    for i, (t, b) in enumerate(steps):
        col = i % 3
        row = i // 3
        _card(s, 0.7 + col * 4.15, 1.35 + row * 2.35, 3.95, 2.15, t, b)
    _footer(s, "04")

    # 5 Architecture
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _add_bg(s)
    _bar(s)
    _title(s, "Architecture — evidence for claims")
    _body(s, [
        "Input → Ingest → Decode → Provenance → Static → Behavioral Twin → Risk → Gate",
        "                                                              ↓",
        "                                        Protected Agent (caps)  |  Analyst Quarantine",
        "",
        "Supporting controls (enterprise-ready):",
        "  • Capability tokens fail-closed on secrets, shell, external email",
        "  • SQLite audit trail + incident JSON export",
        "  • Session risk memory for multi-step jailbreaks",
        "  • Optional OpenAI-compatible LLM planner enrichment (offline twin by default)",
        "",
        "Stack: Python · FastAPI · Streamlit (local) · Vercel production UI",
        "Repo: github.com/krabhi75/aegis-prompt-injection-firewall",
    ], top=1.3, size=15)
    _footer(s, "05")

    # 6 Demo proof
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _add_bg(s)
    _bar(s)
    _title(s, "Demo proof")
    rows = [
        ("Benign product PDF", "ALLOW", "Trusted summarize path"),
        ("Indirect injection in PDF", "BLOCK", "Twin drift: get_secret + send_email"),
        ("Encoded HTML jailbreak", "BLOCK", "Decode layer + override"),
        ("Multi-step roleplay → tools", "QUARANTINE / BLOCK", "Session risk escalation"),
        ("Human analyst queue", "HITL", "Approve / deny with audit note"),
    ]
    y = 1.35
    for name, decision, detail in rows:
        shape = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.7), Inches(y), Inches(11.9), Inches(0.85))
        shape.fill.solid()
        shape.fill.fore_color.rgb = WHITE
        shape.line.color.rgb = LINE
        t = s.shapes.add_textbox(Inches(0.95), Inches(y + 0.18), Inches(5.5), Inches(0.5))
        _set_run(t.text_frame.paragraphs[0].add_run(), name, 14, True, INK)
        d = s.shapes.add_textbox(Inches(6.6), Inches(y + 0.18), Inches(2.2), Inches(0.5))
        color = ALLOW if decision == "ALLOW" else (BLOCK if "BLOCK" in decision else ACCENT)
        _set_run(d.text_frame.paragraphs[0].add_run(), decision, 13, True, color)
        det = s.shapes.add_textbox(Inches(8.8), Inches(y + 0.18), Inches(3.5), Inches(0.5))
        _set_run(det.text_frame.paragraphs[0].add_run(), detail, 12, False, MUTED)
        y += 0.95
    _footer(s, "06")

    # 7 F3/D2
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _add_bg(s)
    _bar(s)
    _title(s, "9-blocker claim — F3 / D2")
    _card(s, 0.7, 1.35, 6.0, 4.6, "F3 — Features (≥7 attack types)",
          "All 9 official attack types implemented and covered in corpus:\n\n"
          "1 Instruction Override\n2 Role Change\n3 Secret Extraction\n"
          "4 Tool Abuse\n5 Credential Theft\n6 Context Poisoning\n"
          "7 Multi-Step Jailbreak\n8 Encoded Instructions\n9 Indirect Prompt Injection\n\n"
          "Live demo demonstrates ≥7 clearly.")
    _card(s, 6.95, 1.35, 5.6, 4.6, "D2 — Depth (high reliability)",
          "Mostly structured/textual multimodal inputs "
          "(user, PDF text, HTML, email, API, OCR text) "
          "with demonstrable reliability:\n\n"
          "• Corpus n = 55\n"
          "• Precision / Recall / F1 ≈ 1.0 on suite\n"
          "• Pytest: 11/11 passed\n"
          "• Command Center metrics in live UI\n\n"
          "Not claiming D3: Word/image OCR not fully shown in production demo.")
    _footer(s, "07")

    # 8 Business + closer
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _add_bg(s)
    _bar(s)
    _title(s, "Business impact & ask")
    _body(s, [
        "Measurable outcomes for enterprises deploying agents:",
        "  • Risk reduction — block secret theft and tool abuse before execution",
        "  • Auditability — incident JSON + decision trail for security teams",
        "  • Governance — OWASP LLM Top 10 & NIST AI RMF mapped controls",
        "  • Human oversight — quarantine band keeps humans in the loop",
        "",
        "Live: https://aegis-prompt-injection-firewall.vercel.app",
        "Code: https://github.com/krabhi75/aegis-prompt-injection-firewall",
        "",
        "Closer:",
        "  Everyone else ships a classifier. We ship a twin.",
        "  If the document tries to make your agent steal secrets,",
        "  the plans diverge — and Aegis blocks before a single tool runs.",
    ], top=1.3, size=15)
    _footer(s, "08")

    prs.save(OUT)
    return OUT


if __name__ == "__main__":
    path = build()
    print(path)
